import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from PIL import Image
import imagehash
import io
import os
import sqlite3
import requests
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

# --- KONFIGURASI ---
st.set_page_config(page_title="Audit Foto Duplicate", layout="wide")
st.title("💾 AUDIT FOTO DUPLICATE")
st.markdown("### Deteksi Otomatis Duplikasi")

# --- DATABASE ENGINE (Menyimpan riwayat foto selamanya) ---
def get_db_connection():
    # Database ini akan menyimpan hash foto agar tidak bisa upload ulang di masa depan
    conn = sqlite3.connect('master_audit_history.db', check_same_thread=False)
    conn.execute('''CREATE TABLE IF NOT EXISTS history 
                 (hash TEXT PRIMARY KEY, segment TEXT, cluster TEXT, info TEXT, date TEXT)''')
    return conn

# --- TURBO DOWNLOADER ---
def smart_download(url):
    if not url or not isinstance(url, str) or "google.com" not in url:
        return None
    try:
        # Deteksi ID dari GDocs atau GDrive
        file_id_match = re.search(r'/(?:d|document/d|file/d)/([a-zA-Z0-9-_]+)', url)
        if not file_id_match: return None
        file_id = file_id_match.group(1)
        
        direct_url = f'https://drive.google.com/uc?export=download&id={file_id}'
        
        # Stream download dengan timeout singkat agar cepat
        response = requests.get(direct_url, timeout=5, stream=True)
        if response.status_code == 200:
            img = Image.open(io.BytesIO(response.content)).convert('RGB')
            img.thumbnail((150, 150)) # Perkecil untuk efisiensi RAM
            return img
    except:
        return None
    return None

# --- CORE AUDIT ENGINE ---
def run_master_audit(file_path):
    wb = load_workbook(file_path, data_only=True)
    results = []
    
    for sheet in wb.worksheets:
        # 1. Deteksi Foto yang MENEMPEL (PNG/JPG) langsung di Excel
        if hasattr(sheet, '_images'):
            for img_obj in sheet._images:
                try:
                    row = img_obj.anchor._from.row + 1
                    col = img_obj.anchor._from.col + 1
                    # Metadata (Kolom B=Cluster, C=Segment)
                    cluster = sheet.cell(row=row, column=2).value or "N/A"
                    segment = sheet.cell(row=row, column=3).value or "N/A"
                    
                    img = Image.open(io.BytesIO(img_obj._data())).convert('RGB')
                    img.thumbnail((150, 150))
                    h = str(imagehash.phash(img))
                    
                    results.append({
                        "Sheet": sheet.title,
                        "Source": "Embedded Foto",
                        "Location": f"Baris {row}, Kolom {col}",
                        "Cluster": cluster,
                        "Segment": segment,
                        "Hash": h,
                        "Img_Obj": img
                    })
                except: continue

        # 2. Deteksi Foto dari LINK (GDocs/GDrive) - Kolom G
        urls_to_process = []
        for r in range(4, sheet.max_row + 1):
            url = sheet.cell(row=r, column=7).value
            if url and "google.com" in str(url):
                urls_to_process.append((r, url, sheet.cell(row=r, column=2).value, sheet.cell(row=r, column=3).value, sheet.title))

        def process_url(item):
            r_idx, url, cluster, segment, s_name = item
            img = smart_download(url)
            if img:
                return {
                    "Sheet": s_name,
                    "Source": "Cloud Link (GDocs/GDrive)",
                    "Location": f"Baris {r_idx}, Kolom G",
                    "Cluster": cluster or "N/A",
                    "Segment": segment or "N/A",
                    "Hash": str(imagehash.phash(img)),
                    "Img_Obj": img
                }
            return None

        # Gunakan 10 pekerja sekaligus agar tidak lama
        with ThreadPoolExecutor(max_workers=10) as executor:
            cloud_res = list(executor.map(process_url, urls_to_process))
            results.extend([res for res in cloud_res if res])

    return results

# --- UI LOGIC ---
uploaded = st.file_uploader("Upload File Patroli EJBN (.xlsx)", type=["xlsx"])

if uploaded:
    if st.button("🚀 MULAI AUDIT LINTAS PERIODE"):
        with st.status("Sedang bekerja keras mengaudit foto...", expanded=True) as status:
            with open("temp_master.xlsx", "wb") as f: f.write(uploaded.getbuffer())
            
            raw_data = run_master_audit("temp_master.xlsx")
            
            if raw_data:
                df = pd.DataFrame(raw_data)
                conn = get_db_connection()
                
                # 1. Cek Duplikat Internal (dalam file yang sama)
                df['Is_Internal_Dup'] = df.duplicated('Hash', keep='first')
                
                # 2. Cek Duplikat Lintas Bulan (History Database)
                history_info = []
                for h in df['Hash']:
                    res = conn.execute("SELECT info, date FROM history WHERE hash=?", (h,)).fetchone()
                    if res:
                        history_info.append(f"⚠️ RE-UPLOAD! (Pernah dipakai di {res[1]})")
                    else:
                        history_info.append("✅ FOTO BARU")
                
                df['History_Status'] = history_info
                
                # Keputusan Akhir
                def final_check(row):
                    if "⚠️" in row['History_Status']: return "❌ GUGUR (Pernah Terbit)"
                    if row['Is_Internal_Dup']: return "❌ GUGUR (Duplikat di Excel Ini)"
                    return "✅ VALID"
                
                df['Status_Akhir'] = df.apply(final_check, axis=1)
                
                # UPDATE DATABASE (Simpan yang VALID ke history agar bulan depan terdeteksi)
                current_date = datetime.now().strftime("%Y-%m-%d")
                for _, r in df[df['Status_Akhir'] == "✅ VALID"].iterrows():
                    try:
                        conn.execute("INSERT INTO history VALUES (?,?,?,?,?)", 
                                     (r['Hash'], str(r['Segment']), str(r['Cluster']), "Audit Verified", current_date))
                    except: pass
                conn.commit()

                status.update(label="Audit Selesai!", state="complete")
                
                # DOWNLOAD SECTION
                st.success(f"Ditemukan {len(df)} Foto. Silakan unduh hasil laporan di bawah.")
                report_df = df.drop(columns=['Img_Obj', 'Hash'])
                towrite = io.BytesIO()
                report_df.to_excel(towrite, index=False)
                st.download_button("📥 DOWNLOAD HASIL AUDIT (EXCEL)", towrite.getvalue(), f"Hasil_Audit_{current_date}.xlsx")

                # TABEL HASIL
                st.dataframe(report_df, use_container_width=True)

                # GALERI
                st.divider()
                st.markdown("#### Preview Audit")
                cols = st.columns(4)
                for idx, r in df.iterrows():
                    with cols[idx % 4]:
                        color = "red" if "❌" in r['Status_Akhir'] else "green"
                        st.image(r['Img_Obj'], caption=f"{r['Location']} | {r['Status_Akhir']}")
            else:
                st.warning("Tidak ditemukan data foto yang valid untuk diaudit.")

    if os.path.exists("temp_master.xlsx"): os.remove("temp_master.xlsx")
